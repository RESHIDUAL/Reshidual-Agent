import os
import re
from typing import List, Dict, Any, Optional
from pathlib import Path
import httpx

class DocumentParser:
    def __init__(self):
        pass

    def extract_text(self, file_path: str) -> str:
        ext = os.path.splitext(file_path)[1].lower()
        if ext == '.pdf':
            return self._extract_pdf(file_path)
        elif ext in ('.docx', '.doc'):
            return self._extract_docx(file_path)
        elif ext in ('.xlsx', '.xls'):
            return self._extract_excel(file_path)
        else:
            return self._extract_plain(file_path)

    def _extract_pdf(self, file_path: str) -> str:
        try:
            from pypdf import PdfReader
            reader = PdfReader(file_path)
            pages = []
            for idx, page in enumerate(reader.pages):
                text = page.extract_text()
                if text:
                    pages.append(f"[Page {idx + 1}]\n{text.strip()}")
            return "\n\n".join(pages)
        except Exception as e:
            return f"Error extracting PDF {file_path}: {e}"

    def _extract_docx(self, file_path: str) -> str:
        try:
            import docx
            doc = docx.Document(file_path)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            return "\n\n".join(paragraphs)
        except Exception as e:
            return f"Error extracting Word document {file_path}: {e}"

    def _extract_excel(self, file_path: str) -> str:
        try:
            import openpyxl
            wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
            sheet_texts = []
            for name in wb.sheetnames:
                sheet = wb[name]
                rows_text = []
                for row in sheet.iter_rows(values_only=True):
                    row_vals = [str(cell) for cell in row if cell is not None and str(cell).strip()]
                    if row_vals:
                        rows_text.append(" | ".join(row_vals))
                if rows_text:
                    sheet_texts.append(f"[Sheet: {name}]\n" + "\n".join(rows_text))
            return "\n\n".join(sheet_texts)
        except Exception as e:
            return f"Error extracting Excel {file_path}: {e}"

    def _extract_plain(self, file_path: str) -> str:
        try:
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                return f.read()
        except Exception as e:
            return f"Error reading text {file_path}: {e}"

    async def extract_url(self, url: str) -> Dict[str, str]:
        headers = {
            "User-Agent": "ReshidualAgent/1.0 (Local-First Context Fetcher)"
        }
        async with httpx.AsyncClient(follow_redirects=True, timeout=15.0) as client:
            resp = await client.get(url, headers=headers)
            resp.raise_for_status()
            html = resp.text

        title_match = re.search(r'<title>(.*?)</title>', html, re.IGNORECASE | re.DOTALL)
        title = title_match.group(1).strip() if title_match else url

        text = re.sub(r'<script[\s\S]*?</script>', '', html, flags=re.IGNORECASE)
        text = re.sub(r'<style[\s\S]*?</style>', '', text, flags=re.IGNORECASE)
        text = re.sub(r'<[^>]+>', ' ', text)
        text = re.sub(r'\s+', ' ', text).strip()

        return {
            "url": url,
            "title": title,
            "content": text[:50000]
        }